"""
Builder script to generate the redesigned, highly presentable, table-centric,
faculty-facing report:
- docs/BEFORE_AFTER_PROJECT_COMPLETE_REPORT_PRESENTABLE.docx
- docs/BEFORE_AFTER_PROJECT_COMPLETE_REPORT_PRESENTABLE.pdf (via MS Word COM)
"""

import os
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = Path(r"C:\Users\PARAS\Desktop\AI SEM 5 B1 G5\PROJECT\cxr-reliability\cxr-reliability")
DOCS_DIR = PROJECT_ROOT / "docs"
DOCS_DIR.mkdir(parents=True, exist_ok=True)

TARGET_DOCX = DOCS_DIR / "BEFORE_AFTER_PROJECT_COMPLETE_REPORT_PRESENTABLE.docx"
TARGET_PDF = DOCS_DIR / "BEFORE_AFTER_PROJECT_COMPLETE_REPORT_PRESENTABLE.pdf"


def set_cell_background(cell, hex_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def set_cell_border(cell, top="E2E8F0", bottom="E2E8F0", left="none", right="none"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    if top != "none":
        borders_xml += f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{top}"/>'
    else:
        borders_xml += '<w:top w:val="none"/>'
    if bottom != "none":
        borders_xml += f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{bottom}"/>'
    else:
        borders_xml += '<w:bottom w:val="none"/>'
    if left != "none":
        borders_xml += f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{left}"/>'
    else:
        borders_xml += '<w:left w:val="none"/>'
    if right != "none":
        borders_xml += f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{right}"/>'
    else:
        borders_xml += '<w:right w:val="none"/>'
    borders_xml += '</w:tcBorders>'
    tc_pr.append(parse_xml(borders_xml))


def add_callout(doc, text: str, title: str = "KEY FINDING", border_hex="0284C7", bg_hex="F0F9FF", icon="📌"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Check current section width
    sec = doc.sections[-1]
    avail_width = sec.page_width - sec.left_margin - sec.right_margin
    tbl.columns[0].width = avail_width
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"{icon} {title}: ")
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(9.5)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    run_body = p.add_run(text)
    run_body.font.name = "Calibri"
    run_body.font.size = Pt(9.5)
    run_body.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(3)


def create_styled_table(doc, headers: list[str], rows: list[list[str]], col_widths: list[float], align_right_cols: list[int] = None):
    """
    col_widths: list of inches (e.g. [1.5, 2.5, ...])
    align_right_cols: list of column indices to right-align (e.g. numeric columns)
    """
    if align_right_cols is None:
        align_right_cols = []
        
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Header row
    hdr_row = tbl.rows[0]
    tr_pr = hdr_row._tr.get_or_add_trPr()
    tr_pr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))  # Repeat header row across pages
    
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = Inches(col_widths[i])
        set_cell_background(cell, "1E293B")  # Dark slate
        set_cell_margins(cell, top=75, bottom=75, left=90, right=90)
        set_cell_border(cell, top="1E293B", bottom="0F172A", left="none", right="none")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i in align_right_cols else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(h)
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    # Data rows
    for r_idx, r_data in enumerate(rows):
        row = tbl.rows[r_idx + 1]
        bg = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.width = Inches(col_widths[c_idx])
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=55, bottom=55, left=90, right=90)
            set_cell_border(cell, top="E2E8F0", bottom="E2E8F0", left="none", right="none")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if c_idx in align_right_cols else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(str(val))
            run.font.name = "Calibri"
            run.font.size = Pt(8.0)
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(2)
    return tbl


def add_h1(doc, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)  # Navy
    return p


def add_h2(doc, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)  # Primary blue
    return p


def add_h3(doc, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)  # Slate
    return p


def add_body_p(doc, text: str, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p


def add_qa_item(doc, q: str, m: str, res_headers: list[str], res_rows: list[list[str]], res_widths: list[float], interp: str, align_right: list[int] = None):
    """
    Standard visual layout for each experimental result:
    QUESTION -> METHOD -> RESULT TABLE -> INTERPRETATION
    """
    add_h3(doc, f"QUESTION: {q}")
    add_body_p(doc, f"Method: {m}")
    create_styled_table(doc, res_headers, res_rows, res_widths, align_right_cols=align_right)
    add_callout(doc, interp, title="INTERPRETATION", border_hex="10B981", bg_hex="F0FDF4", icon="💡")


def build_presentable_report():
    print("Building presentable document...")
    doc = docx.Document()
    
    # Configure initial portrait margins
    sec0 = doc.sections[0]
    sec0.top_margin = Inches(0.8)
    sec0.bottom_margin = Inches(0.8)
    sec0.left_margin = Inches(0.8)
    sec0.right_margin = Inches(0.8)
    sec0.page_width = Inches(8.5)
    sec0.page_height = Inches(11.0)
    
    # -------------------------------------------------------------
    # COVER PAGE
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(54)
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run("BEFORE vs AFTER")
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(26)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    r_s = p_sub.add_run("Project Improvement & Experimental Evaluation")
    r_s.font.name = "Calibri"
    r_s.font.size = Pt(16)
    r_s.font.bold = True
    r_s.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
    
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(0)
    p_desc.paragraph_format.space_after = Pt(28)
    r_d = p_desc.add_run("Reliability-Aware Multi-Agent System for Chest X-Ray Pneumonia Classification")
    r_d.font.name = "Calibri"
    r_d.font.size = Pt(12)
    r_d.font.italic = True
    r_d.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    meta_headers = ["Metadata Field", "Authoritative Specification"]
    meta_rows = [
        ["Dataset", "NIH ChestX-ray14"],
        ["Frozen Test Cohort", "16,724 frontal radiographs (4,621 patients)"],
        ["Diagnostic Model", "TorchXRayVision DenseNet-121 (densenet121-res224-nih)"],
        ["Architecture", "7-Agent Reliability Supervisory Pipeline"],
        ["Operating Threshold", "0.522161 (Calibrated F1-Optimal Operating Point)"],
        ["Project Classification", "Academic Research Prototype"],
        ["Evaluation Date", "October 6, 2026"]
    ]
    create_styled_table(doc, meta_headers, meta_rows, [2.3, 4.6])
    
    add_callout(
        doc,
        "BEFORE: 'What is the model prediction?' → AFTER: 'Is the prediction reliable enough to release?' The system introduces an active supervisory reliability layer around a frozen DenseNet-121 classifier to govern prediction release, repair, and human clinical review.",
        title="CORE PARADIGM SHIFT",
        border_hex="0284C7",
        bg_hex="F0F9FF"
    )
    
    doc.add_page_break()
    
    # -------------------------------------------------------------
    # SECTION 1 — PROJECT AT A GLANCE
    # -------------------------------------------------------------
    add_h1(doc, "Section 1 — Project at a Glance")
    
    add_h2(doc, "System Topology Comparison")
    
    topo_headers = ["Baseline Pipeline (BEFORE)", "Reliability Multi-Agent Pipeline (AFTER)"]
    topo_rows = [
        [
            "Input Chest Radiograph\n↓\nDenseNet-121\n↓\nUnconditional Prediction\n(Forced on 100% of images)",
            "Input Chest Radiograph\n↓\nQuality Assessment (Laplacian, SNR, Exposure)\n↓\nDenseNet-121 Classification + 1024-D Feature Extraction\n↓\nProbability Calibration + Uncertainty + OOD Screening\n↓\nDecision Arbitration (Accept / Repair / Escalate / Reject)\n↓\nTargeted Repair (CLAHE / NLMeans / Unsharp)\n↓\nFresh Post-Repair Inference Pass\n↓\nMulti-Gate Verification (5 Safety Gates)\n↓\nSelective Release OR Triage to Radiologist Review"
        ]
    ]
    create_styled_table(doc, topo_headers, topo_rows, [3.2, 3.7])
    
    add_h2(doc, "Core Architectural Differences")
    arch_headers = ["System Capability", "Conventional Baseline (BEFORE)", "Reliability System (AFTER)"]
    arch_rows = [
        ["Image Quality Screening", "None (blind to degradation)", "Quality Agent (Laplacian, SNR, Histogram)"],
        ["Out-of-Distribution Screening", "None (silent outlier failure)", "Mahalanobis Distance via Cholesky (1024-D)"],
        ["Uncertainty Representation", "None (raw sigmoid output)", "Platt Calibration, Confidence, Shannon Entropy"],
        ["Targeted Defect Repair", "None (no restoration)", "Repair Agent (CLAHE, NLMeans, Unsharp Mask)"],
        ["Post-Repair Validation", "None (unverified modification)", "Verification Agent (5 mandatory safety gates)"],
        ["Selective Prediction Release", "No (100% unconditional release)", "Yes (43.94% coverage on high-trust cases)"],
        ["Human Review Escalation", "No (unassisted autonomous predictor)", "Yes (56.06% safely triaged to radiologist)"]
    ]
    create_styled_table(doc, arch_headers, arch_rows, [2.0, 2.3, 2.6])
    
    add_callout(
        doc,
        "BEFORE: 'What is the model prediction?' | AFTER: 'Is the prediction reliable enough to release?'",
        title="CENTRAL RESEARCH QUESTION",
        border_hex="0284C7",
        bg_hex="F0F9FF",
        icon="🎯"
    )
    
    # -------------------------------------------------------------
    # SECTION 2 — EXECUTIVE RESULTS
    # -------------------------------------------------------------
    add_h1(doc, "Section 2 — Executive Results")
    
    exec_headers = ["Project Evaluation Area", "Empirical Measurement", "Operational Status"]
    exec_rows = [
        ["Full Frozen Test Cohort", "16,724 images", "Authoritative benchmark"],
        ["Baseline Automated Coverage", "100.0% (16,724 images)", "Unconditional release"],
        ["Reliability Automated Released", "7,349 images", "High-trust cohort"],
        ["Reliability Coverage Rate", "43.94%", "Selective release"],
        ["Withheld for Human Review", "9,375 images", "Safety triage"],
        ["Withheld Percentage", "56.06%", "Intentional safety mechanism"],
        ["Baseline Total Diagnostic Errors", "1,515 errors (1,344 FP + 171 FN)", "Unsupervised error release"],
        ["Errors among Released Cases", "565 errors (498 FP + 67 FN)", "High-certainty release"],
        ["Baseline Errors Withheld (Contained)", "950 errors (846 FP + 104 FN)", "62.71% error containment"],
        ["Quality Repairs Evaluated", "12,082 repairs", "Conditional execution"],
        ["Repair Quality Recovery Rate", "78.28% (9,458 / 12,082)", "Objective signal gain"],
        ["Severe OOD Cases Rejected", "101 / 101 (100.0%)", "Complete outlier interception"],
        ["Borderline OOD Cases Escalated", "86 / 86 (100.0%)", "Complete borderline triage"],
        ["Paired Diagnostic Validation Cohort", "35 repaired cases", "Authoritative paired trial"],
        ["Diagnostic Errors Corrected by Repair", "0 errors", "Zero correction gain"],
        ["New Diagnostic Errors Introduced", "0 errors", "Zero harmful degradation"],
        ["DenseNet Prediction Stability", "100.0% (0 prediction flips)", "Complete stability"]
    ]
    create_styled_table(doc, exec_headers, exec_rows, [2.7, 2.3, 1.9], align_right_cols=[1])
    
    add_callout(
        doc,
        "950 diagnostic errors present in the standalone baseline were not present among automatically released cases because the corresponding cases were safely withheld from automated release. This represents error containment through selective release, not ground-truth correction during inference.",
        title="WHAT THIS MEANS — SCIENTIFIC ERROR CONTAINMENT",
        border_hex="10B981",
        bg_hex="F0FDF4",
        icon="🛡️"
    )
    
    # -------------------------------------------------------------
    # SECTION 3 — BASE MODEL BEFORE VS AFTER
    # -------------------------------------------------------------
    add_h1(doc, "Section 3 — Base Model Before vs After")
    
    bm_headers = ["Diagnostic Metric", "Standalone DenseNet (BEFORE)", "Reliability Released (AFTER)", "Statistical Interpretation"]
    bm_rows = [
        ["Evaluated Population", "Full test cohort (N=16,724)", "Released subset (N=7,349)", "Fundamentally distinct populations"],
        ["Automated Coverage", "100.0% (0 withheld)", "43.94% (9,375 withheld)", "Selective gating of low-confidence inputs"],
        ["Classification Accuracy", "90.9412%", "92.3119%", "Released subset accuracy (+1.37%)"],
        ["Precision (PPV)", "3.5176%", "3.3010%", "Reflects extreme class imbalance (1.32% prev)"],
        ["Recall (Sensitivity)", "22.2727%", "20.2381%", "Frozen classifier sensitivity"],
        ["Specificity (TNR)", "91.8565%", "93.1452%", "Slight TN purity gain in released set"],
        ["Negative Predictive Value", "98.8846%", "99.0196%", "Exceptionally high exclusion confidence"],
        ["F1-Score", "0.060756", "0.056761", "Heavily constrained by 1.32% prevalence"],
        ["Area Under ROC (AUROC)", "0.701619", "—", "Baseline ranking discrimination"],
        ["Area Under PR Curve (AUPRC)", "0.028085", "—", "Baseline precision-recall curve"],
        ["Total Released Errors", "1,515 errors", "565 errors", "950 errors safely contained"]
    ]
    create_styled_table(doc, bm_headers, bm_rows, [1.8, 1.7, 1.7, 1.7], align_right_cols=[1, 2])
    
    add_callout(
        doc,
        "92.31% accuracy is NOT presented as a direct improvement in the underlying DenseNet classifier because the denominator changed (7,349 vs 16,724). It is the accuracy of the selectively released subset. The classifier weights remained frozen throughout the project.",
        title="IMPORTANT STATISTICAL INTERPRETATION",
        border_hex="DC2626",
        bg_hex="FEF2F2",
        icon="⚠️"
    )
    
    doc.add_page_break()
    
    # -------------------------------------------------------------
    # SECTION 4 — AGENT-WISE RESULTS (UNIFORM STRUCTURE FOR ALL 7 AGENTS)
    # -------------------------------------------------------------
    add_h1(doc, "Section 4 — Agent-Wise Evaluation")
    add_body_p(doc, "Each agent is evaluated using a uniform template: Purpose, Method, Selection Rationale, Quantitative Results, and Actual Improvement.")
    
    # --- AGENT 1: QUALITY AGENT ---
    add_h2(doc, "Agent 1 — Quality Agent")
    add_body_p(doc, "Purpose: Detect physical signal and acquisition defects (blur, noise, exposure) before predictions are trusted.", italic=True)
    add_body_p(doc, "Method: Computes Laplacian variance for sharpness, SNR (dB) for noise, and mean intensity/histogram statistics for exposure dynamic range.")
    add_body_p(doc, "Why this method? Fast O(N) convolution requiring zero additional neural network inference; deterministic, robust, and directly interpretable.")
    
    q_state_headers = ["Quality Classification", "Radiograph Count", "Cohort Percentage", "Operational Action"]
    q_state_rows = [
        ["GOOD", "3,608", "21.57%", "Eligible for direct evaluation"],
        ["DEGRADED", "940", "5.62%", "Requires repair or escalation"],
        ["POOR", "12,176", "72.81%", "Routed to targeted repair"]
    ]
    create_styled_table(doc, q_state_headers, q_state_rows, [2.0, 1.6, 1.6, 1.7], align_right_cols=[1, 2])
    
    q_rec_headers = ["Repair Transition State", "Evaluated Count", "Transition Rate", "Quality Gate Outcome"]
    q_rec_rows = [
        ["POOR → GOOD", "7,766", "64.28%", "Full recovery (Passes Gate 4)"],
        ["POOR → DEGRADED", "1,692", "14.00%", "Partial recovery (Blocked by Gate 4)"],
        ["POOR → POOR", "2,624", "21.72%", "Unresolved defect (Blocked by Gate 4)"],
        ["Total Improved", "9,458", "78.28%", "Objective physical signal gain"],
        ["Total Worsened", "0", "0.00%", "Zero quality degradation observed"]
    ]
    create_styled_table(doc, q_rec_headers, q_rec_rows, [2.0, 1.6, 1.6, 1.7], align_right_cols=[1, 2])
    
    add_callout(
        doc,
        "What it actually improved: Added explicit objective image-quality assessment. Identified 12,176 poor-quality cases and achieved 78.28% quality improvement across 12,082 evaluated repairs. (Note: This is objective signal recovery, NOT termed 'accuracy' due to lack of ground-truth quality labels).",
        title="AGENT 1 CONTRIBUTION",
        border_hex="0284C7",
        bg_hex="F0F9FF"
    )
    
    # --- AGENT 2: DENSENET BASE MODEL ---
    add_h2(doc, "Agent 2 — DenseNet-121 Base Model")
    add_body_p(doc, "Purpose: Serves as the sole disease classification engine and extracts a 1024-dimensional semantic feature representation.", italic=True)
    add_body_p(doc, "Method: Evaluates TorchXRayVision DenseNet-121 (densenet121-res224-nih) at calibrated operating threshold tau = 0.522161.")
    add_body_p(doc, "Why this method? Pretrained on multi-institutional chest radiography archives; dense feature reuse captures lung opacities while extracting GAP features.")
    
    cm_headers = ["Confusion Matrix (Full Test)", "Actual Negative (16,504)", "Actual Pneumonia (220)"]
    cm_rows = [
        ["Predicted Negative (Non-Pneumonia)", "15,160 (True Negative)", "171 (False Negative)"],
        ["Predicted Positive (Pneumonia)", "1,344 (False Positive)", "49 (True Positive)"]
    ]
    create_styled_table(doc, cm_headers, cm_rows, [2.7, 2.1, 2.1], align_right_cols=[1, 2])
    
    add_callout(
        doc,
        "What it actually improved: Maintained stable diagnostic classification across all tests. Provided the disease prediction and rich 1024-dimensional semantic embeddings. Its standalone error count (1,515 errors) remained unchanged as weights were frozen.",
        title="AGENT 2 CONTRIBUTION",
        border_hex="0284C7",
        bg_hex="F0F9FF"
    )
    
    # --- AGENT 3: CALIBRATION & UNCERTAINTY AGENT ---
    add_h2(doc, "Agent 3 — Calibration & Uncertainty Agent")
    add_body_p(doc, "Purpose: Maps raw network scores to calibrated probabilities and quantifies predictive confidence and Shannon entropy.", italic=True)
    add_body_p(doc, "Method: Platt Scaling (logistic regression on validation logits), confidence = max(p, 1-p), and normalized binary Shannon entropy.", italic=False)
    add_body_p(doc, "Why this method? Deep networks produce uncalibrated overconfident scores; Platt scaling preserves ROC ranking while minimizing calibration error.")
    
    cal_headers = ["Probability Metric", "Raw Sigmoid Output", "Platt Calibrated", "Empirical Delta"]
    cal_rows = [
        ["Expected Calibration Error (ECE)", "~0.051", "~0.015", "-0.036 (Superior probability alignment)"],
        ["Brier Score", "0.0128", "0.0121", "-0.0007 (Reduced quadratic error)"]
    ]
    create_styled_table(doc, cal_headers, cal_rows, [2.3, 1.5, 1.5, 1.6], align_right_cols=[1, 2, 3])
    
    unc_headers = ["Uncertainty Tier", "Cohort Count", "Cohort %", "Released Subset", "Released Accuracy"]
    unc_rows = [
        ["LOW Uncertainty", "2,653", "15.86%", "1,642 (61.89%)", "99.57% (1,635 TN, 0 TP, 0 FP, 7 FN)"],
        ["HIGH Uncertainty", "14,071", "84.14%", "5,707 (40.56%)", "90.22% (5,132 TN, 17 TP, 498 FP, 60 FN)"]
    ]
    create_styled_table(doc, unc_headers, unc_rows, [1.8, 1.3, 1.1, 1.3, 1.4], align_right_cols=[1, 2, 3, 4])
    
    add_callout(
        doc,
        "What it actually improved: Probability trustworthiness and risk stratification. LOW uncertainty predictions achieved 99.57% accuracy in the released set (+9.35% over high uncertainty). High uncertainty does not mean incorrect, but identifies radiographs near the boundary.",
        title="AGENT 3 CONTRIBUTION",
        border_hex="0284C7",
        bg_hex="F0F9FF"
    )
    
    # --- AGENT 4: OOD AGENT ---
    add_h2(doc, "Agent 4 — Out-of-Distribution (OOD) Detection Agent")
    add_body_p(doc, "Purpose: Detect inputs whose 1024-dimensional feature representations diverge from the in-distribution training manifold.", italic=True)
    add_body_p(doc, "Method: Mahalanobis distance in 1024-D feature space, solved via Cholesky decomposition of the covariance matrix with Tikhonov regularization.")
    add_body_p(doc, "Why this method? Euclidean distance ignores feature correlations; Mahalanobis distance scales by directional feature covariance; Cholesky ensures numerical stability.")
    
    ood_headers = ["Distribution Status", "Distance Threshold", "Count", "Percentage", "Routing Enforcement"]
    ood_rows = [
        ["IN-DISTRIBUTION", "D_M <= 35.54 (<= 95th pct)", "16,537", "98.88%", "Continue in pipeline"],
        ["BORDERLINE", "35.54 < D_M <= 42.65 (95-99th pct)", "86", "0.51%", "100.0% Escalated to Review (86/86)"],
        ["SEVERE OUTLIER", "D_M > 42.65 (> 99th pct)", "101", "0.60%", "100.0% Intercepted & Rejected (101/101)"]
    ]
    create_styled_table(doc, ood_headers, ood_rows, [1.8, 1.8, 1.1, 1.1, 1.1], align_right_cols=[2, 3])
    
    add_callout(
        doc,
        "What it actually improved: Distributional screening and safety enforcement. 100% of severe outliers (101/101) were rejected; 100% of borderline cases (86/86) were escalated. (Formal OOD AUROC requires an external non-chest benchmark and is not claimed on NIH alone).",
        title="AGENT 4 CONTRIBUTION",
        border_hex="0284C7",
        bg_hex="F0F9FF"
    )
    
    # --- AGENT 5: DECISION AGENT ---
    add_h2(doc, "Agent 5 — Decision Agent")
    add_body_p(doc, "Purpose: Deterministically arbitrates multi-agent evidence into clinical actions (ACCEPT, REPAIR, ESCALATE, REJECT).", italic=True)
    add_body_p(doc, "Method: Hierarchical rule engine enforcing PRD FR-5 rules R1–R7 with strict precedence.")
    add_body_p(doc, "Why this method? In clinical AI safety, supervisory arbitration must be deterministic, transparent, and 100% auditable; avoids black-box meta-classifiers.")
    
    dec_init_headers = ["Initial Routing Action", "Trigger Condition", "Case Count", "Percentage"]
    dec_init_rows = [
        ["ACCEPT (Direct)", "Quality GOOD, OOD In-Dist, Uncertainty LOW", "501", "3.00%"],
        ["REPAIR", "Quality POOR/DEGRADED, OOD In-Dist", "13,001", "77.74%"],
        ["ESCALATE (Direct)", "Quality GOOD with High Uncertainty, or Borderline OOD", "3,121", "18.66%"],
        ["REJECT", "OOD SEVERE Outlier", "101", "0.60%"],
        ["Total Initial Decisions", "Complete test cohort reconciliation", "16,724", "100.00%"]
    ]
    create_styled_table(doc, dec_init_headers, dec_init_rows, [1.8, 2.5, 1.3, 1.3], align_right_cols=[2, 3])
    
    dec_final_headers = ["Final Operational Disposition", "Contributing Pathways", "Case Count", "Percentage"]
    dec_final_rows = [
        ["Direct Release", "Initial ACCEPT routing", "501", "3.00%"],
        ["Released After Repair", "REPAIR → Fresh Pass → Verification PASS", "6,848", "40.94%"],
        ["TOTAL AUTOMATIC RELEASE", "Direct Release + Verified Repair Release", "7,349", "43.94%"],
        ["WITHHELD FOR HUMAN REVIEW", "Direct Escalate (3,121) + Repair Escalate (5,234) + Reject (101) + Bounds (919)", "9,375", "56.06%"]
    ]
    create_styled_table(doc, dec_final_headers, dec_final_rows, [2.0, 2.5, 1.2, 1.2], align_right_cols=[2, 3])
    
    add_callout(
        doc,
        "Initial routing describes the first decision on incoming images. Final disposition describes what happened after repair and verification. They are fundamentally distinct operational concepts.",
        title="ROUTING vs DISPOSITION DISTINCTION",
        border_hex="0284C7",
        bg_hex="F0F9FF"
    )
    
    # --- AGENT 6: REPAIR AGENT ---
    add_h2(doc, "Agent 6 — Image Repair Agent")
    add_body_p(doc, "Purpose: Applies targeted, non-destructive restoration to in-distribution radiographs exhibiting technical defects.", italic=True)
    add_body_p(doc, "Method: Unsharp Masking for blur; Fast Non-Local Means for noise; CLAHE for exposure.")
    add_body_p(doc, "Why this method? Targeted algorithms address specific physical degradation modes rather than applying a blunt universal filter.")
    
    rep_stat_headers = ["Repair Execution Metric", "Measured Value", "Operational Interpretation"]
    rep_stat_rows = [
        ["Repair Agent Invocations", "12,176 cases", "Defect identified by Quality Agent"],
        ["Repairs Physically Applied", "12,082 cases", "Restoration executed on array"],
        ["Repairs Skipped", "94 cases", "Within acceptable tolerance margins"],
        ["Repairs Refused", "0 cases", "Zero bounding exceptions"],
        ["Quality Metric Improvement Rate", "78.28% (9,458 / 12,082)", "Objective signal gain"],
        ["Quality Metric Worsening Rate", "0.00% (0 / 12,082)", "Zero quality degradation"],
        ["Diagnostic Errors Corrected", "0 errors (N=35 Paired Validation)", "Zero accuracy gain for frozen CNN"]
    ]
    create_styled_table(doc, rep_stat_headers, rep_stat_rows, [2.5, 2.0, 2.4], align_right_cols=[1])
    
    add_callout(
        doc,
        "Repair improved measurable image quality (78.28% recovery), but the project did NOT demonstrate that repair improves pneumonia diagnostic correctness for the frozen DenseNet-121.",
        title="CRITICAL REPAIR AGENT FINDING",
        border_hex="DC2626",
        bg_hex="FEF2F2",
        icon="⚖️"
    )
    
    # --- AGENT 7: VERIFICATION AGENT ---
    add_h2(doc, "Agent 7 — Verification Agent")
    add_body_p(doc, "Purpose: Audits post-repair radiographs across five mandatory safety gates before authorizing automated release.", italic=True)
    add_body_p(doc, "Method: Evaluates repair execution, label stability, OOD status, quality recovery to GOOD, and confidence non-degradation.")
    add_body_p(doc, "Why this method? Prevents failed repairs or repair-induced hallucinations from reaching clinical workflow.")
    
    ver_stat_headers = ["Verification Outcome", "Repaired Cases", "Percentage", "Operational Routing"]
    ver_stat_rows = [
        ["Verified & Released (PASS)", "6,848", "56.68%", "Released as ACCEPTED_AFTER_REPAIR"],
        ["Escalated to Human Review (FAIL)", "5,234", "43.32%", "Withheld and flagged for radiologist"],
        ["Total Repairs Evaluated", "12,082", "100.00%", "Complete verification coverage"]
    ]
    create_styled_table(doc, ver_stat_headers, ver_stat_rows, [2.3, 1.5, 1.5, 1.6], align_right_cols=[1, 2])
    
    ver_gate_headers = ["Safety Gate", "Evaluation Criterion", "Threshold Guard", "Purpose"]
    ver_gate_rows = [
        ["Gate 1: Execution", "Repair applied", "Array modified == True", "Verify physical operation"],
        ["Gate 2: Label Stability", "Diagnostic prediction", "No flip across tau = 0.522161", "Prevent hallucination"],
        ["Gate 3: Distribution", "OOD feature check", "Mahalanobis D_M <= 35.54", "Ensure in-distribution"],
        ["Gate 4: Quality Recovery", "Post-repair status", "Must transition to GOOD", "Ensure technical recovery"],
        ["Gate 5: Non-Degradation", "Confidence & SNR", "Delta_conf >= -0.01, Delta_SNR >= -5 dB", "Prevent confidence collapse"]
    ]
    create_styled_table(doc, ver_gate_headers, ver_gate_rows, [1.5, 1.5, 2.0, 1.9])
    
    doc.add_page_break()
    
    # -------------------------------------------------------------
    # SECTION 5 — REPAIR PARAMETER TUNING
    # -------------------------------------------------------------
    add_h1(doc, "Section 5 — Repair Parameter Tuning Results")
    add_body_p(doc, "The project systematically evaluated 29 distinct algorithmic configurations across 580 inference runs on a patient-isolated validation cohort (N=60, 0% test set overlap).")
    
    tune_headers = ["Degradation Mode", "Configuration Tested", "Image Quality Effect", "Model Stability", "Operational Decision"]
    tune_rows = [
        ["Blur (Unsharp)", "r=1.0, amount=0.50 (Prod)", "Var 58.31 → 116.02 (+57.71)", "100.0% (0 flips)", "RETAIN in production"],
        ["Blur (Unsharp)", "r=2.0, amount=0.50", "Var 58.31 → 127.00 (+68.69)", "100.0% (0 flips)", "No diagnostic advantage"],
        ["Blur (Unsharp)", "r=3.0, amount=1.00", "Var 58.31 → 220.41 (+162.1)", "100.0% (conf -2.8%)", "Edge ringing penalty"],
        ["Noise (NLMeans)", "h=3.0 (Prod)", "SNR gain = 0.00 dB (18.96 dB)", "100.0% (0 flips)", "Weak smoothing weight"],
        ["Noise (NLMeans)", "h=7.0", "SNR gain = +0.57 dB (19.53 dB)", "100.0% (0 flips)", "Conservative candidate"],
        ["Noise (NLMeans)", "h=10.0", "SNR gain = +6.60 dB (25.57 dB)", "95.0% (1 harmful flip)", "Risk: Erases infiltrate"],
        ["Noise (NLMeans)", "h=15.0", "SNR gain = +19.91 dB (38.87 dB)", "95.0% (1 harmful flip)", "Excessive over-smoothing"],
        ["Exposure (CLAHE)", "clip=1.0, grid=8x8", "PSNR 21.91 dB, MAE 18.66", "90.0% (2 flips)", "High radiometric fidelity"],
        ["Exposure (CLAHE)", "clip=2.0, grid=8x8 (Prod)", "PSNR 20.35 dB, MAE 20.42", "85.0% (3 flips)", "RETAIN; monitor drift"],
        ["Exposure (CLAHE)", "clip=4.0, grid=8x8", "PSNR 17.74 dB, MAE 26.00", "85.0% (3 flips)", "Noise amplification"]
    ]
    create_styled_table(doc, tune_headers, tune_rows, [1.4, 1.6, 1.6, 1.2, 1.1])
    
    add_callout(
        doc,
        "Across all 29 configurations and 580 inference runs, no net diagnostic improvement was demonstrated. Repair is an image-quality restoration mechanism, not a proven diagnostic enhancement mechanism for the frozen DenseNet.",
        title="TUNING EXPERIMENT CONCLUSION",
        border_hex="0284C7",
        bg_hex="F0F9FF"
    )
    
    # -------------------------------------------------------------
    # SECTION 6 — FINAL PAIRED DIAGNOSTIC VALIDATION (N=35)
    # -------------------------------------------------------------
    add_h1(doc, "Section 6 — Final Paired Repair → DenseNet Diagnostic Validation")
    add_body_p(doc, "Authoritative validation resolving: 'After the Repair Agent repairs a poor-quality X-ray, does passing the repaired X-ray through DenseNet-121 again improve pneumonia classification?'")
    
    p_trans_headers = ["Diagnostic Transition", "Repaired Cases (N=35)", "Percentage", "Clinical Interpretation"]
    p_trans_rows = [
        ["Correct → Correct", "17", "48.57%", "True prediction preserved"],
        ["Correct → Incorrect", "0", "0.00%", "Zero harmful prediction flips"],
        ["Incorrect → Correct", "0", "0.00%", "Zero diagnostic errors corrected"],
        ["Incorrect → Incorrect", "18", "51.43%", "Diagnostic error persisted"],
        ["Total Repaired Cohort", "35", "100.00%", "Complete prediction stability"]
    ]
    create_styled_table(doc, p_trans_headers, p_trans_rows, [2.0, 1.5, 1.5, 1.9], align_right_cols=[1, 2])
    
    p_met_headers = ["Diagnostic Metric", "Before Repair", "After Repair", "Metric Delta", "Interpretation"]
    p_met_rows = [
        ["Classification Accuracy", "48.5714%", "48.5714%", "+0.0000%", "Unchanged accuracy"],
        ["Precision (PPV)", "44.4444%", "44.4444%", "+0.0000%", "Unchanged precision"],
        ["Recall (Sensitivity)", "23.5294%", "23.5294%", "+0.0000%", "Unchanged sensitivity"],
        ["Specificity (TNR)", "72.2222%", "72.2222%", "+0.0000%", "Unchanged specificity"],
        ["F1-Score", "0.307692", "0.307692", "+0.000000", "Unchanged balance"],
        ["Negative Predictive Value", "50.0000%", "50.0000%", "+0.0000%", "Unchanged NPV"],
        ["Area Under ROC (AUROC)", "0.483660", "0.480392", "-0.003268", "Minimal continuous drift"],
        ["Area Under PR Curve", "0.482922", "0.479935", "-0.002987", "Minimal continuous drift"]
    ]
    create_styled_table(doc, p_met_headers, p_met_rows, [1.8, 1.2, 1.2, 1.2, 1.5], align_right_cols=[1, 2, 3])
    
    add_callout(
        doc,
        "DIAGNOSTIC ERRORS CORRECTED: 0 | NEW DIAGNOSTIC ERRORS: 0 | PREDICTION STABILITY: 100.0%. Repair improved image quality in 85.71% of cases, but the frozen DenseNet-121 prediction did not change diagnostically.",
        title="AUTHORITATIVE PAIRED EXPERIMENT RESULT",
        border_hex="10B981",
        bg_hex="F0FDF4",
        icon="🏆"
    )
    
    doc.add_page_break()
    
    # -------------------------------------------------------------
    # SECTION 7 — REPAIR BY DEFECT COHORT
    # -------------------------------------------------------------
    add_h1(doc, "Section 7 — Repair by Defect Cohort")
    
    add_h2(doc, "A. Blur Cohort (Unsharp Masking, N=20 Paired Cases)")
    b_headers = ["Metric", "Before Repair", "After Repair", "Change", "Gate Status"]
    b_rows = [
        ["Laplacian Variance", "58.31", "116.02", "+57.71 (+98.97%)", "Passed (>100 threshold)"],
        ["Accuracy (tau=0.522161)", "60.00%", "60.00%", "0.00%", "100% stable"],
        ["Prediction Flips", "—", "—", "0 flips", "12 Corr→Corr, 8 Inc→Inc"],
        ["Errors Corrected", "—", "—", "0 errors", "Zero correction gain"]
    ]
    create_styled_table(doc, b_headers, b_rows, [1.8, 1.3, 1.3, 1.4, 1.1], align_right_cols=[1, 2, 3])
    
    add_h2(doc, "B. Noise Cohort (Fast Non-Local Means)")
    n_headers = ["Evaluation Mode", "Pre-Repair SNR", "Post-Repair SNR", "SNR Delta", "Diagnostic Effect"]
    n_rows = [
        ["Natural Noise (Paired N=6)", "40.12 dB", "40.72 dB", "+0.60 dB", "1 Corr→Corr, 5 Inc→Inc, 0 flips"],
        ["Synthetic Denoising h=3.0", "18.96 dB", "18.96 dB", "+0.00 dB", "0 flips (weak kernel weight)"],
        ["Synthetic Denoising h=7.0", "18.96 dB", "19.53 dB", "+0.57 dB", "0 flips (effective candidate)"],
        ["Synthetic Denoising h=10.0", "18.96 dB", "25.57 dB", "+6.60 dB", "1 harmful flip (Correct→Incorrect)"],
        ["Synthetic Denoising h=15.0", "18.96 dB", "38.87 dB", "+19.91 dB", "1 harmful flip (over-smoothing)"]
    ]
    create_styled_table(doc, n_headers, n_rows, [1.8, 1.2, 1.2, 1.2, 1.5], align_right_cols=[1, 2, 3])
    
    add_h2(doc, "C. Exposure Cohort (CLAHE)")
    e_headers = ["Evaluation Mode", "Shifted Intensity", "Repaired Intensity", "Delta", "Diagnostic Effect"]
    e_rows = [
        ["Natural Exposure (Paired N=9)", "112.45", "121.80", "+9.35", "4 Corr→Corr, 5 Inc→Inc, 0 flips"],
        ["Synthetic CLAHE clip=1.0", "115.42", "123.01", "+7.59", "1 helpful flip, 1 harmful flip"],
        ["Synthetic CLAHE clip=2.0 (Prod)", "115.42", "123.95", "+8.53", "1 helpful flip, 2 harmful flips"],
        ["Synthetic CLAHE clip=4.0", "115.42", "124.62", "+9.20", "Noise amplification in soft tissue"]
    ]
    create_styled_table(doc, e_headers, e_rows, [1.8, 1.2, 1.2, 1.2, 1.5], align_right_cols=[1, 2, 3])
    
    # -------------------------------------------------------------
    # SECTION 8 — DELIBERATELY BLURRED STRESS TEST
    # -------------------------------------------------------------
    add_h1(doc, "Section 8 — Deliberately Blurred Stress Test: Repair & Verification Trace")
    add_body_p(doc, "A chest radiograph with severe artificial motion blur was fed into the end-to-end pipeline to verify that the Repair Agent does not blindly force acceptance and that the Verification Agent halts unrecovered images.")
    
    st_trace_headers = ["Pipeline Stage", "Evaluated Parameter", "Measured Value", "Operational Significance"]
    st_trace_rows = [
        ["Pre-Repair Screening", "Quality Status", "POOR", "Severe blur defect identified"],
        ["Pre-Repair Screening", "Laplacian Variance", "1.70", "Far below 100.0 sharpness threshold"],
        ["Pre-Repair Screening", "Signal-to-Noise Ratio", "59.4 dB", "High SNR; defect is purely spatial"],
        ["Pre-Repair Screening", "DenseNet Raw Score", "0.5041", "Boundary prediction (tau=0.522161)"],
        ["Pre-Repair Screening", "Uncertainty / OOD", "HIGH / In-Distribution (27.98)", "Eligible for Rule R3 Repair"],
        ["Initial Routing", "Decision Agent Action", "REPAIR", "Correctly dispatched to Repair Agent"],
        ["Repair Execution", "Algorithm & Parameters", "Unsharp Mask (r=1.0, amount=0.5)", "Array modified == TRUE"],
        ["Post-Repair Inference", "Laplacian Variance", "1.72 (Delta = +0.016)", "Minimal high-frequency gain"],
        ["Post-Repair Inference", "Post-Quality Status", "POOR", "Failed to recover to GOOD (< 100)"],
        ["Post-Repair Inference", "DenseNet Raw Score", "0.504083 (Delta = -0.000003)", "Zero diagnostic classification flip"],
        ["Verification Gate 1", "Execution Audit", "PASS", "Array physically altered"],
        ["Verification Gate 2", "Label Stability", "PASS", "No prediction flip"],
        ["Verification Gate 3", "Distribution Check", "PASS", "OOD remains In-Distribution (27.99)"],
        ["Verification Gate 4", "Quality Recovery", "FAIL (POOR → POOR)", "Sharpness 1.72 << 100.0 threshold"],
        ["Verification Gate 5", "Non-Degradation Guard", "PASS", "Delta_conf = -0.0003 >= -0.01"],
        ["FINAL DISPOSITION", "Verification Outcome", "ESCALATE TO RADIOLOGIST", "Blocked from automated release"]
    ]
    create_styled_table(doc, st_trace_headers, st_trace_rows, [1.6, 1.8, 1.8, 1.7])
    
    add_callout(
        doc,
        "REPAIR DID NOT FORCE ACCEPTANCE. The repair was correctly executed, but the severe blur had destroyed high frequencies that unsharp masking cannot hallucinate. Because quality remained POOR (1.72 vs 100), Verification correctly escalated the case to human review.",
        title="STRESS TEST SAFETY CONCLUSION",
        border_hex="0284C7",
        bg_hex="F0F9FF",
        icon="🛡️"
    )
    
    doc.add_page_break()
    
    # -------------------------------------------------------------
    # SECTION 9 — SELECTIVE RELEASE & ERROR CONTAINMENT
    # -------------------------------------------------------------
    add_h1(doc, "Section 9 — Selective Release & Error Containment")
    
    add_h2(doc, "Selective Release Distribution")
    sr_headers = ["Disposition Pathway", "Radiograph Count", "Percentage", "Visual Coverage Representation"]
    sr_rows = [
        ["Directly Accepted (High Quality, Low Uncertainty)", "501", "3.00%", "[■■                    ]"],
        ["Accepted After Verified Repair", "6,848", "40.94%", "[■■■■■■■■■■            ]"],
        ["TOTAL AUTOMATIC RELEASE", "7,349", "43.94%", "[■■■■■■■■■■■           ] 43.94%"],
        ["WITHHELD FOR RADIOLOGIST REVIEW", "9,375", "56.06%", "[■■■■■■■■■■■■■■        ] 56.06%"]
    ]
    create_styled_table(doc, sr_headers, sr_rows, [2.5, 1.2, 1.0, 2.2], align_right_cols=[1, 2])
    
    add_h2(doc, "Error Containment Analysis (Baseline vs Released)")
    ec_headers = ["Error Classification", "Baseline Errors (N=16,724)", "Released Errors (N=7,349)", "Errors Withheld (Contained)", "Containment Rate"]
    ec_rows = [
        ["False Positives (FP)", "1,344", "498", "846", "62.95%"],
        ["False Negatives (FN)", "171", "67", "104", "60.82%"],
        ["TOTAL DIAGNOSTIC ERRORS", "1,515", "565", "950", "62.71%"]
    ]
    create_styled_table(doc, ec_headers, ec_rows, [2.1, 1.2, 1.2, 1.2, 1.2], align_right_cols=[1, 2, 3, 4])
    
    add_callout(
        doc,
        "Errors were contained from automated release through selective withholding, NOT through ground-truth correction during inference. Withholding 56.06% of cases prevented 950 baseline diagnostic errors from entering clinical workflow.",
        title="ERROR CONTAINMENT DEFINITION",
        border_hex="10B981",
        bg_hex="F0FDF4",
        icon="✅"
    )
    
    # -------------------------------------------------------------
    # SECTION 10 — OVERALL BEFORE VS AFTER MASTER TABLE (LANDSCAPE)
    # -------------------------------------------------------------
    doc.add_page_break()
    sec_land1 = doc.add_section()
    sec_land1.orientation = WD_ORIENT.LANDSCAPE
    sec_land1.page_width = Inches(11.0)
    sec_land1.page_height = Inches(8.5)
    sec_land1.top_margin = Inches(0.7)
    sec_land1.bottom_margin = Inches(0.7)
    sec_land1.left_margin = Inches(0.7)
    sec_land1.right_margin = Inches(0.7)
    
    add_h1(doc, "Section 10 — Overall Before vs After Comprehensive Comparison")
    add_body_p(doc, "Master multidimensional comparison between the conventional baseline pipeline and the reliability-aware multi-agent supervisory framework.")
    
    master_comp_headers = ["Dimension", "Conventional Baseline (BEFORE)", "Reliability Multi-Agent (AFTER)", "Empirical Evidence", "Clinical Interpretation"]
    master_comp_rows = [
        ["Architecture", "Monolithic feedforward CNN", "7-Agent supervisory pipeline", "End-to-end Python/PyTorch codebase", "Modular, auditable clinical workflow"],
        ["Quality Assessment", "Zero inspection (blind)", "Laplacian, SNR, Intensity screening", "12,176 poor cases flagged on test set", "Prevents garbage-in, garbage-out"],
        ["Blur Handling", "Processed unconditionally", "Laplacian < 100 flags; Unsharp mask", "Tuning: +57.71 variance; Paired: 100% stab", "Restores edges before reliance"],
        ["Noise Handling", "Passed directly to CNN", "SNR < 15 dB flags; NLMeans applied", "Tuning: h=7 to 10 gave +0.6 to +6.6 dB SNR", "Attenuates sensor noise floor"],
        ["Exposure Handling", "Processed unconditionally", "Mean outside [20, 235] flags; CLAHE", "Intensity recovered 115.42 → 123.95", "Prevents dynamic range clipping"],
        ["Disease Classification", "DenseNet-121 (tau=0.522161)", "DenseNet-121 (identical frozen weights)", "Zero weight or architecture modifications", "Diagnostic engine remains identical"],
        ["Probability Calibration", "Raw sigmoid scores (ECE ~ 0.051)", "Platt calibrated probabilities", "ECE reduced to ~0.015 on validation", "Reflects empirical event rates"],
        ["Uncertainty Routing", "Unrepresented", "Confidence + Shannon Entropy", "Low uncertainty = 99.57% accuracy", "Actionable triage thresholding"],
        ["OOD Detection", "Unmonitored", "1024-D Mahalanobis via Cholesky", "100% severe rejected (101), 100% bord esc (86)", "Prevents silent outlier failures"],
        ["Decision Routing", "100% prediction release", "Deterministic R1–R7 rule engine", "501 Acc, 13,001 Rep, 3,121 Esc, 101 Rej", "Transparent, auditable actions"],
        ["Targeted Repair", "Absent", "CLAHE, NLMeans, Unsharp Masking", "12,082 repairs evaluated across test set", "Salvages repairable radiographs"],
        ["Post-Repair Verification", "Absent", "5-gate safety firewall", "6,848 releases, 5,234 escalations", "Intercepts unrecovered images"],
        ["Human Clinical Review", "Absent (autonomous predictor)", "Formal escalation pathway", "9,375 cases safely triaged to radiologist", "Realistic human-AI partnership"],
        ["Automated Coverage", "100.0% (16,724 released)", "43.94% (7,349 released)", "Selective release on high-trust cases", "Automates only validated inputs"],
        ["Diagnostic Error Release", "1,515 errors released", "565 errors released", "950 baseline errors withheld", "62.71% error containment"],
        ["Quality Recovery Rate", "0% (unmodified)", "78.28% (9,458 / 12,082)", "Poor→Good: 7,766; Poor→Degraded: 1,692", "Strong physical signal recovery"],
        ["Repair Diagnostic Gain", "Not applicable", "0 errors corrected in N=35", "17 Corr→Corr, 18 Inc→Inc, 0 flips", "Repair fixes quality, not diagnosis"]
    ]
    create_styled_table(doc, master_comp_headers, master_comp_rows, [1.4, 2.0, 2.2, 2.0, 2.0])
    
    # -------------------------------------------------------------
    # SECTION 11 — WHAT ACTUALLY IMPROVED (EVIDENCE VS CONCLUSION)
    # -------------------------------------------------------------
    doc.add_page_break()
    add_h1(doc, "Section 11 — What Actually Improved: Evidence vs Scientific Conclusion")
    
    what_headers = ["Project Evaluation Area", "Experimental Evidence Generated", "Scientific Finding", "Status"]
    what_rows = [
        ["Image Quality Assessment", "Identified 12,176 poor cases out of 16,724", "Quantifies blur, noise, and exposure objectively", "DEMONSTRATED"],
        ["Quality Recovery", "78.28% of repaired cases improved (9,458 / 12,082)", "Restores measurable physical signal properties", "DEMONSTRATED"],
        ["OOD Safety Routing", "101/101 severe rejected, 86/86 borderline escalated", "Screening prevents silent failures on outliers", "DEMONSTRATED"],
        ["Uncertainty Stratification", "Low uncertainty released accuracy = 99.57%", "Confidence and entropy identify reliable subsets", "DEMONSTRATED"],
        ["Decision Determinism", "16,724 cases perfectly reconciled across R1–R7", "Rule table provides 100% auditable triage", "DEMONSTRATED"],
        ["Targeted Repair Execution", "12,082 repairs physically executed on uint8 arrays", "Modular filters execute conditionally", "DEMONSTRATED"],
        ["Diagnostic Gain from Repair", "0 errors corrected across N=35 paired benchmark", "Repair does not improve frozen CNN accuracy", "NOT DEMONSTRATED"],
        ["Verification Gatekeeping", "Blocked 5,234 repairs; escalated severe blur test", "Prevents unrecovered images from releasing", "DEMONSTRATED"],
        ["Selective Release", "43.94% automated coverage (7,349 cases)", "Limits automated release to high-trust cases", "DEMONSTRATED"],
        ["Error Containment", "950 baseline errors withheld (62.71% containment)", "Substantially purifies released predictions", "DEMONSTRATED"],
        ["Clinical Safety / Replacement", "Research prototype; no external radiologist trial", "Does not replace thoracic radiologists", "REQUIRES VALIDATION"]
    ]
    create_styled_table(doc, what_headers, what_rows, [1.7, 2.8, 3.2, 1.9])
    
    # -------------------------------------------------------------
    # SECTION 12 — MASTER METHOD SELECTION RATIONALE (LANDSCAPE)
    # -------------------------------------------------------------
    doc.add_page_break()
    add_h1(doc, "Section 12 — Consolidated Master Method Selection Rationale")
    
    master_meth_headers = ["Agent / Function", "Algorithmic Method", "Why Selected", "What it Measures", "Primary Known Limitation"]
    master_meth_rows = [
        ["Quality: Sharpness", "Laplacian Variance", "O(N) convolution; highly sensitive to edge loss", "High-frequency spatial 2nd derivative", "Invariant to blur direction; noise sensitive"],
        ["Quality: Noise", "Signal-to-Noise Ratio (dB)", "Direct signal dominance over background floor", "10*log10(mu^2 / sigma_noise^2)", "Background patch estimation variance"],
        ["Quality: Exposure", "Mean & Histogram Stats", "Lightweight bounds on sensor saturation", "Global dynamic range across [0, 255]", "Cannot detect localized focal underexposure"],
        ["Base Model: Classifier", "DenseNet-121 (nih)", "Dense connectivity; CXR pretraining; 1024-D GAP", "Pneumonia probability & latent features", "Frozen model; vulnerable to corrupted input"],
        ["Calibration: Posterior", "Platt Scaling (Logistic)", "Monotonic; preserves ROC ranking; lightweight", "Calibrated posterior probability p_cal", "Assumes sigmoid miscalibration; no AUC gain"],
        ["Uncertainty: Dispersion", "Confidence & Shannon Entropy", "Quantifies boundary ambiguity non-linearly", "max(p, 1-p) and normalized entropy in [0, 1]", "Does not separate aleatoric from epistemic"],
        ["OOD: Feature Shift", "Mahalanobis Distance (1024-D)", "Accounts for latent covariance and correlation", "Statistical divergence from training manifold", "Assumes unimodal Gaussian feature distribution"],
        ["OOD: Numerical Solver", "Cholesky Decomposition", "Numerically stable; avoids explicit 1024x1024 inversion", "Solves L * y = (x - mu), distance = ||y||_2", "Requires Tikhonov epsilon regularization"],
        ["Decision: Arbiter", "Hierarchical Rule Engine (R1–R7)", "100% deterministic; auditable; zero black-box risk", "Maps multi-agent signals to clinical action", "Static rules do not dynamically adapt"],
        ["Repair: Blur Defect", "Unsharp Masking (r=1, a=0.5)", "Lightweight; restores edges without ringing", "High-frequency local gradient boosting", "Cannot reconstruct destroyed spatial detail"],
        ["Repair: Noise Defect", "Fast Non-Local Means", "Exploits anatomical patch self-similarity", "Weighted non-local patch neighborhood average", "Over-smoothing risk (h >= 10 erases opacities)"],
        ["Repair: Exposure Defect", "CLAHE (clip=2.0, grid=8x8)", "Restores local contrast; limits noise amplification", "Local adaptive histogram remapping", "Non-physical remapping; causes CNN drift"],
        ["Verification: Safety Gate", "5-Gate Safety Architecture", "Defense-in-depth; ensures post-repair integrity", "Execution, stability, OOD, recovery, confidence", "Conservative gating withholds benign cases"]
    ]
    create_styled_table(doc, master_meth_headers, master_meth_rows, [1.5, 1.8, 2.2, 2.2, 1.9])
    
    # -------------------------------------------------------------
    # SECTION 13 — FACULTY SUMMARY (ONE-PAGE EXECUTIVE PORTRAIT)
    # -------------------------------------------------------------
    doc.add_page_break()
    sec_port1 = doc.add_section()
    sec_port1.orientation = WD_ORIENT.PORTRAIT
    sec_port1.page_width = Inches(8.5)
    sec_port1.page_height = Inches(11.0)
    sec_port1.top_margin = Inches(0.8)
    sec_port1.bottom_margin = Inches(0.8)
    sec_port1.left_margin = Inches(0.8)
    sec_port1.right_margin = Inches(0.8)
    
    add_h1(doc, "Section 13 — Faculty Executive Summary: 3-Box Synthesis")
    
    add_callout(
        doc,
        "A conventional CAD system based on TorchXRayVision DenseNet-121 that takes an input chest radiograph and unconditionally produces a pneumonia prediction regardless of severe blur, extreme sensor noise, demographic shift, or model uncertainty. Evaluated on the full NIH test set (N=16,724), it forced 1,515 baseline diagnostic errors into clinical workflow.",
        title="BOX 1: WHAT WE STARTED WITH (BASELINE)",
        border_hex="475569",
        bg_hex="F8FAFC",
        icon="📦"
    )
    
    add_callout(
        doc,
        "An active seven-agent supervisory reliability architecture wrapping the frozen DenseNet-121 classifier: Quality Agent (physical screening), Base Model (classification & 1024-D features), Calibration & Uncertainty Agent (Platt scaling & entropy), OOD Agent (Mahalanobis distance via Cholesky), Decision Agent (deterministic R1–R7 rule table), Repair Agent (targeted CLAHE/NLM/Unsharp), and Verification Agent (5 post-repair safety gates).",
        title="BOX 2: WHAT WE BUILT (MULTI-AGENT ARCHITECTURE)",
        border_hex="0284C7",
        bg_hex="F0F9FF",
        icon="⚙️"
    )
    
    add_callout(
        doc,
        "1. Image quality screening and recovery (78.28% improvement across 12,082 repairs).\n2. Distributional safety screening (100% severe OOD rejected, 100% borderline escalated).\n3. Error containment: 950 baseline errors withheld (62.71% containment rate; released accuracy 92.31%).\n4. Verified safety gatekeeping: Severe blur stress test correctly escalated after repair failed to recover quality.\n5. Scientific honesty: In the authoritative paired N=35 experiment, repair corrected 0 diagnostic errors for the frozen DenseNet. Repair restores image quality, NOT frozen model accuracy.",
        title="BOX 3: WHAT WE DEMONSTRATED (EMPIRICAL FINDINGS)",
        border_hex="10B981",
        bg_hex="F0FDF4",
        icon="📊"
    )
    
    # -------------------------------------------------------------
    # SECTION 14 — FACULTY DEFENSE Q&A (CONCISE & STRUCTURED)
    # -------------------------------------------------------------
    add_h1(doc, "Section 14 — Faculty Oral Defense Q&A: Key Technical Answers")
    
    qa_list = [
        ("What was the project before?", "A conventional pipeline feeding radiographs directly into DenseNet-121, forcing predictions on 100% of images and releasing 1,515 baseline errors."),
        ("What did you add?", "A 7-agent supervisory reliability layer: Quality, Calibration/Uncertainty, OOD, Decision, Repair, and Verification agents, plus a selective release and human review architecture."),
        ("What is the actual contribution?", "Shifting medical AI from unconditional 'prediction-only AI' to 'reliability-aware AI' that knows when to predict, repair, verify, escalate, or withhold."),
        ("Why DenseNet-121?", "TorchXRayVision's checkpoint is pretrained on vast chest radiograph archives; dense connectivity captures pulmonary opacities while extracting 1024-D GAP features."),
        ("Why Laplacian Variance?", "Lightweight O(N) second spatial derivative measuring edge sharpness without neural network overhead."),
        ("Why Signal-to-Noise Ratio?", "Directly quantifies anatomical signal dominance over sensor noise floor in decibels."),
        ("Why Mean Intensity / Histogram?", "Fast, deterministic bounding on global radiometric dynamic range ([20, 235]) to prevent clipping."),
        ("Why Platt Scaling?", "Logistic calibration on validation logits minimizes Expected Calibration Error (0.051 → 0.015) while preserving ROC ranking."),
        ("Why Confidence & Entropy?", "max(p, 1-p) and normalized Shannon entropy mathematically bound predictive dispersion in [0, 1], stratifying risk (99.57% low vs 90.22% high)."),
        ("Why Mahalanobis & Cholesky?", "Mahalanobis accounts for latent feature covariance across 1024 channels; Cholesky solves the linear system stably without explicit matrix inversion."),
        ("Why Rule-Based Decision Agent?", "Safety arbitration in healthcare must be deterministic, auditable, and transparent, eliminating black-box meta-classifier risks."),
        ("Why CLAHE, NLMeans, and Unsharp?", "Targeted algorithms: CLAHE redistributes local contrast; NLMeans exploits anatomical patch self-similarity; Unsharp boosts local edge gradients."),
        ("Did repair improve diagnostic accuracy?", "No. In the authoritative paired experiment of 35 repaired cases, diagnostic correctness was unchanged (0 errors corrected, 100% stability). Repair fixes image quality, not frozen CNN accuracy."),
        ("What happened in the severe blur stress test?", "A radiograph with severe blur (variance 1.70 vs 100) was repaired via Unsharp Mask. Variance reached only 1.72 (POOR → POOR). Verification failed Gate 4 and correctly escalated to human review, proving Repair != Automatic Acceptance."),
        ("How many cases were released vs withheld?", "7,349 cases (43.94%) were automatically released (501 direct + 6,848 verified repairs); 9,375 cases (56.06%) were safely withheld for radiologist review."),
        ("How did the system contain errors?", "By withholding unverified and degraded radiographs, 950 baseline errors (846 FP + 104 FN) were contained from automated release (62.71% error containment rate).")
    ]
    
    for q_text, a_text in qa_list:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.line_spacing = 1.15
        r_q = p_q.add_run(f"Q: {q_text} ")
        r_q.font.name = "Calibri"
        r_q.font.size = Pt(9.5)
        r_q.font.bold = True
        r_q.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
        
        r_a = p_q.add_run(f"A: {a_text}")
        r_a.font.name = "Calibri"
        r_a.font.size = Pt(9.5)
        r_a.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        
    doc.add_page_break()
    
    # -------------------------------------------------------------
    # FINAL PAGE — VISUAL PARADIGM SUMMARY
    # -------------------------------------------------------------
    p_final_t = doc.add_paragraph()
    p_final_t.paragraph_format.space_before = Pt(72)
    p_final_t.paragraph_format.space_after = Pt(12)
    p_final_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ft = p_final_t.add_run("THE PARADIGM SHIFT")
    r_ft.font.name = "Calibri"
    r_ft.font.size = Pt(22)
    r_ft.font.bold = True
    r_ft.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    final_headers = ["BEFORE: Prediction-Only AI", "AFTER: Reliability-Aware AI"]
    final_rows = [
        [
            "Input Radiograph\n↓\nDenseNet-121\n↓\nForced Prediction\n\n(Releases 1,515 baseline errors)",
            "Input Radiograph\n↓\nAssess Quality\n↓\nPredict & Extract 1024-D Features\n↓\nEvaluate Uncertainty & OOD\n↓\nArbitrate Action (Accept / Repair / Escalate / Reject)\n↓\nRepair Targeted Defect (if indicated)\n↓\nFresh Post-Repair Inference\n↓\nVerify Across 5 Safety Gates\n↓\nRelease Verified OR Triage to Radiologist\n\n(Contains 950 baseline errors)"
        ]
    ]
    create_styled_table(doc, final_headers, final_rows, [3.2, 3.7])
    
    p_concl = doc.add_paragraph()
    p_concl.paragraph_format.space_before = Pt(36)
    p_concl.paragraph_format.space_after = Pt(12)
    p_concl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_concl.add_run("“From prediction-only AI to reliability-aware AI.”")
    r_c.font.name = "Calibri"
    r_c.font.size = Pt(16)
    r_c.font.bold = True
    r_c.font.italic = True
    r_c.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
    
    # Save the redesigned Word docx
    doc.save(str(TARGET_DOCX))
    print(f"Successfully saved redesigned DOCX: {TARGET_DOCX}")
    
    # Convert to PDF via MS Word COM
    print("Converting DOCX to PDF via Microsoft Word COM...")
    try:
        import win32com.client
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc_obj = word.Documents.Open(str(TARGET_DOCX))
        doc_obj.SaveAs(str(TARGET_PDF), FileFormat=17)  # 17 = wdFormatPDF
        doc_obj.Close()
        word.Quit()
        print(f"Successfully generated PDF: {TARGET_PDF}")
    except Exception as e:
        print(f"PDF conversion encountered an issue: {e}")


if __name__ == "__main__":
    build_presentable_report()
